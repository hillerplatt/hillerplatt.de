"""Erzeuge die Webdaten aus der unveränderten Excel-Quelldatei."""

import csv
import json
import os
import re

import openpyxl


PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
WORTLISTE_DIR = os.path.join(PROJECT_ROOT, "wortliste")
TABELLE_PATH = os.path.join(WORTLISTE_DIR, "Tabelle.xlsx")
RECORDER_TOOL_DIR = os.path.join(WORTLISTE_DIR, "RecorderTool")
AUDIO_RECORDER = os.path.join(PROJECT_ROOT, "audio", "recorder")

GERMAN_SORT_TRANS = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue"})


def sort_key(value):
    """Plattformunabhängige Sortierung mit ä/ö/ü wie ae/oe/ue."""
    folded = (value or "").casefold()
    return folded.translate(GERMAN_SORT_TRANS), folded


def clean_whitespace(value):
    if not isinstance(value, str):
        return value
    return "\n".join(re.sub(r"[ \t]+", " ", line).strip()
                     for line in value.splitlines()).strip()


def clean_content_web(value):
    """Erste Excel-Zeile; dieselbe Schreibweise bestimmt Anzeige und Audionamen."""
    if value is None:
        return ""
    lines = str(value).splitlines()
    value = lines[0] if lines else ""
    value = clean_whitespace(value).replace("\\-", "").replace("...", "…")
    return value.replace("?", "︖").replace(":", "：").replace("'", "’").replace("`", "’")


def full_word(entry):
    return f"{entry['artikel']} {entry['plattdeutsch']}" if entry["artikel"] else entry["plattdeutsch"]


def existing_prefixes():
    path = os.path.join(WORTLISTE_DIR, "wortliste.json")
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as file:
        return {(entry["artikel"], entry["plattdeutsch"]): entry["sprecher"]
                for entry in json.load(file)}


def get_prefix(entry, previous):
    name = full_word(entry)
    for prefix in ("ab", "hw"):
        if os.path.isfile(os.path.join(AUDIO_RECORDER, f"{prefix}-{name}.flac")):
            return prefix
    return previous.get((entry["artikel"], entry["plattdeutsch"]), "hw")


def validate_identity(value, location):
    """Ein versehentliches Leerzeichen darf keinen neuen Audiodateinamen erzeugen."""
    if not isinstance(value, str) or not value:
        return
    first_line = value.splitlines()[0]
    if first_line != clean_whitespace(first_line) or re.search(r"\(\s|\s\)", first_line):
        raise ValueError(f"Unsaubere Schreibweise in {location}: {first_line!r}")
    displayed = clean_content_web(first_line)
    if re.search(r'[<>:"/\\|?*\x00-\x1f]', displayed):
        raise ValueError(f"Ungültiges Zeichen im Recorder-Dateinamen in {location}: {displayed!r}")
    if len(displayed.encode("utf-8")) > 240:
        raise ValueError(f"Recorder-Dateiname zu lang in {location}: {displayed!r}")


def build_data():
    workbook = openpyxl.load_workbook(TABELLE_PATH, read_only=True, data_only=True)
    previous = existing_prefixes()
    words = []
    sheet = workbook["Wörterbuch"]
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not row[1] and not row[2]:
            continue
        validate_identity(row[0], f"Wörterbuch!A{row_number}")
        validate_identity(row[1], f"Wörterbuch!B{row_number}")
        entry = {
            "artikel": clean_content_web(row[0]),
            "plattdeutsch": clean_content_web(row[1]),
            "hochdeutsch": clean_content_web(row[2]),
        }
        entry["sprecher"] = get_prefix(entry, previous)
        entry["audio"] = f"{entry['sprecher']}-{full_word(entry)}.flac"
        words.append(entry)
    words.sort(key=lambda entry: sort_key(entry["plattdeutsch"]))

    sayings = []
    sheet = workbook["Redewendungen u. Sprachgebrauch"]
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not row[0]:
            continue
        validate_identity(row[0], f"Redewendungen!A{row_number}")
        plattdeutsch = clean_content_web(row[0])
        sayings.append({
            "plattdeutsch": plattdeutsch,
            "hochdeutsch": clean_content_web(row[1]),
            "audio": f"{plattdeutsch}.flac",
        })
    sayings.sort(key=lambda entry: sort_key(entry["plattdeutsch"]))
    workbook.close()
    return words, sayings


def write_data(words, sayings):
    with open(os.path.join(WORTLISTE_DIR, "wortliste.json"), "w", encoding="utf-8") as file:
        json.dump(words, file, ensure_ascii=False, indent=2)
    with open(os.path.join(WORTLISTE_DIR, "wortliste.js"), "w", encoding="utf-8") as file:
        file.write("const wortliste = ")
        json.dump(words, file, ensure_ascii=False)
    with open(os.path.join(WORTLISTE_DIR, "wortliste.csv"), "w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["artikel", "plattdeutsch", "hochdeutsch", "sprecher", "audio"], lineterminator="\n")
        writer.writeheader()
        writer.writerows(words)
    with open(os.path.join(WORTLISTE_DIR, "redewendungen.json"), "w", encoding="utf-8") as file:
        json.dump(sayings, file, ensure_ascii=False, indent=4)
    with open(os.path.join(WORTLISTE_DIR, "redewendungen.js"), "w", encoding="utf-8") as file:
        file.write("const redewendungen = ")
        json.dump(sayings, file, ensure_ascii=False, indent=4)

    os.makedirs(RECORDER_TOOL_DIR, exist_ok=True)
    with open(os.path.join(RECORDER_TOOL_DIR, "Woerterliste.txt"), "w", encoding="utf-8") as file:
        for entry in sorted(words, key=lambda entry: sort_key(full_word(entry))):
            file.write(full_word(entry) + "\n")
    with open(os.path.join(RECORDER_TOOL_DIR, "Redewendungen.txt"), "w", encoding="utf-8") as file:
        for entry in sayings:
            file.write(entry["plattdeutsch"] + "\n")


if __name__ == "__main__":
    words, sayings = build_data()
    write_data(words, sayings)
    print(f"Erzeugt: {len(words)} Wörter und {len(sayings)} Redewendungen aus {TABELLE_PATH}")
